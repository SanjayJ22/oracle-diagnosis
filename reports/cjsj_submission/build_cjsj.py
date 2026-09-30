"""Build JillellamudiSanjay_paper.docx (CJSJ). Every number traces to
reports/cjsj_numbers.md, reports/research_audit.md or reports/pilot_results.md."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

FONT = "Times New Roman"
BODY = 10
doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.85)
    s.top_margin = s.bottom_margin = Inches(0.7)
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY)
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
st.paragraph_format.space_after = Pt(4)

def para(text="", size=BODY, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         after=4, before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    return p

def rich(parts, size=BODY, after=4):
    """parts: list of (text, 'b'|'i'|'') tuples."""
    p = para(after=after)
    for t, f in parts:
        r = p.add_run(t); r.font.size = Pt(size)
        r.bold = "b" in f; r.italic = "i" in f
    return p

def h1(t): para(t, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3)
def h2(t): para(t, size=BODY, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2)

# ---------------- front matter ----------------
para("Distinguishing Price-Oracle Attacks from Genuine Market Volatility in "
     "Decentralized Finance Using Market Data", size=14, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("Sanjay Jillellamudi", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Prosper High School, Prosper, Texas, United States", size=9.5, italic=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

rich([("Abstract. ", "b"),
      ("Decentralized finance (DeFi) protocols depend on price oracles, and a sudden price anomaly can "
       "signal either an attack or a genuine market move that requires no intervention. This study asked "
       "whether market-observable signals can distinguish the two, and how to tell whether an apparent "
       "answer is real. On 162 Ethereum incidents (24 exploits and 138 large market moves), a Bayesian "
       "network, a random forest, and XGBoost all reached 95 to 98% leave-one-out accuracy. A validity "
       "audit showed that this separation is explained by shortcuts: features describing the overall Ethereum "
       "market, which carry no information about the attacked asset, predicted the label with an AUC of "
       "0.91 (permutation p ≤ 0.001), because the two classes were collected and measured differently; once the shortcuts were removed, no feature carried separately identifiable signal. "
       "The shortcut had a second cost. When models were trained on earlier incidents, including mislabeled ones, and tested on later "
       "ones, XGBoost raised 99 false alarms across seven time splits against 14 for a Bayesian network "
       "whose fixed, domain-set thresholds left the shortcut features nearly inert (paired difference 85, "
       "95% CI 40 to 129), and still raised 48 when trained on verified labels only; removing the shortcut features largely eliminated the bursts. Verifying labels "
       "against two independent post-mortems reduced the attack class from 24 to 9, and a pilot of a "
       "corrected measurement protocol found only 3 of those 9 observable at block resolution (5 with intra-block replay). Detectors "
       "in this domain must therefore be evaluated on anomalies that are sourced and measured identically "
       "across classes; a protocol for doing so is proposed.", "")], size=9.5, after=6)

# ---------------- 1. Introduction ----------------
h1("1. Introduction")
para("Blockchains cannot observe the outside world, so DeFi protocols import asset prices through oracles "
     "such as Chainlink [1]. Lending markets and stablecoins use these prices to value collateral and "
     "trigger liquidations [2]. When a reported price suddenly moves, two very different causes are "
     "possible. An attacker may have distorted it, often with capital borrowed for a single transaction "
     "through a flash loan [3], in which case the protocol should halt. Or the market may have genuinely "
     "moved, in which case halting freezes users' funds for no reason. The 24 exploits examined in this "
     "study alone caused losses of about $370 million.")
para("Most prior work detects attacks from inside the blockchain, using transaction execution traces "
     "(POMABuster [4], DeFiRanger [5]) or event logs and traces (DeFiTrace [6]). TOAD-ML [7] scores oracle "
     "trustworthiness from multi-source market data. "
     "Bayesian networks have been used to separate deliberate attacks from accidental failures in physical "
     "infrastructure [8]. This study asks a complementary question: once an anomaly has occurred, can "
     "signals observable from outside the attacked contract tell an attack from a genuine market move?")
para("A second question proved more important. Machine learning models often succeed by exploiting "
     "shortcuts: features that correlate with the label in a particular dataset but are unrelated to the "
     "phenomenon being studied [9]. This paper shows that a DeFi incident benchmark assembled in the natural "
     "way contains such shortcuts, quantifies them, and shows how models differ in relying on them.")

# ---------------- 2. Methods ----------------
h1("2. Methods")
rich([("Data. ", "b"),
      ("Twenty-four exploits (2020 to 2026) were taken from the DeFiHackLabs corpus [10], each confirmed "
       "on-chain. The 138 negative examples were moments when the Chainlink ETH/USD feed moved by 3 to 5% "
       "or more, for 162 incidents in total. Each incident is described by about a dozen features in two groups. ", ""),
      ("Venue features", "i"),
      (" are measured at the pool or contract involved: reserve skew (percentage change over about 300 "
       "blocks), its localization relative to a basket of reference pools, and price deviations against "
       "other venues. ", ""),
      ("Market-context features", "i"),
      (" describe the Ethereum market: oracle update freshness, network congestion, realized volatility, "
       "and four statistics of the oracle's trailing 24-hour trajectory. All windows are strictly trailing.", "")])
rich([("Models. ", "b"),
      ("The Bayesian network (BN) has a hand-specified structure in which the cause points to every feature, "
       "plus one edge from congestion to update freshness. Features are split at fixed, domain-motivated "
       "thresholds (5% for price-type features, 3,600 s for freshness, and Ethereum's 50% gas target for "
       "congestion), and conditional probability tables are learned with BDeu priors in pgmpy [11]. "
       "Baselines are Gaussian naive Bayes, logistic regression, a random forest, and XGBoost [12] (the "
       "model family used by TOAD-ML; nested-CV tuned for LOOCV; for time splits a fixed configuration with 0.8 row and column subsampling).", "")])
rich([("Evaluation. ", "b"),
      ("Models were scored by leave-one-out cross-validation (LOOCV), by leaving out whole date clusters (149 clusters; identical to LOOCV), "
       "and by forward-in-time splits (train on incidents before a cutoff, test on those after; seven "
       "cutoffs from July 2022 to July 2025). Paired comparisons use exact McNemar tests with Holm "
       "correction; intervals are bootstrap 95% CIs; permutation tests use at least 1,000 label shuffles. "
       "Attributions use SHAP [13] for tree models and log-likelihood evidence for the BN.", "")])
rich([("Shortcut audit and label verification. ", "b"),
      ("Models were trained on feature subsets that cannot describe the attacked asset (market context "
       "only), on missingness indicators, and on the incident date. A subset that predicts the label "
       "despite carrying no information about the attack indicates a shortcut. Each exploit was then "
       "classified by a written rule (an attacker moved a price source that a separate contract read, and "
       "the profit depended on that read), and counted as verified only with two independent post-mortems.", "")])
rich([("Pilot of a corrected protocol. ", "b"),
      ("Both classes entered through the same trigger (the price source itself moved more than 0.5% for "
       "pegged assets or 5% otherwise within 50 blocks), and features were measured on that source as prices "
       "rather than balances. Benign episodes came from documented events on attacked pool types, such as the March 2023 USDC depeg.", "")])

# ---------------- 3. Results ----------------
h1("3. Results")
h2("3.1 High apparent accuracy")
para("Under LOOCV the random forest caught 20 of 24 attacks with no false alarms (97.5% accuracy), XGBoost "
     "caught 21 with none (98.1%), and the BN caught 20 with 4 false alarms among 138 market moves (95.1%). The BN "
     "and random forest were not significantly different (McNemar p = 0.125; p = 0.688 on verified labels).")
h2("3.2 The benchmark contains shortcuts")
para("Figure 1 shows how the classes were assembled. Every market move was chosen because ETH moved and was "
     "measured on one ETH pool, while for 23 of 24 attacks ETH/USD data "
     "stood in as a fallback. A random forest using only market-context features reached an AUC of 0.906 [0.845, 0.960], and 0.917 "
     "[0.802, 0.989] on the 9 verified attacks, matching the venue features (0.901 and 0.808; permutation "
     "p ≤ 0.001 on all 24, ≤ 0.002 throughout; Figure 2). The shortcut is not specific to that model: logistic regression, XGBoost and naive Bayes reached 0.85, "
     "0.92 and 0.83 on those features alone. Two further shortcuts appeared. Three price features were missing "
     "for 83 to 96% of attacks but for no market moves; logistic regression reached 99.4% with missingness "
     "indicators and 88.9% without. Reserve skew was a pool price "
     "change for market moves but, for many attacks, the victim contract's token balance after the exploit: "
     "16 of 24 attack values exceeded 90%, largely the loss itself, while no market move exceeded 19%. The date alone did not predict the label "
     "(AUC 0.575, p = 0.23). With market context, missingness indicators and reserve skew all removed, no "
     "clean signal remained: the four remaining price features were missing for 20 to 23 of 24 attacks, so four "
     "missingness flags alone reached AUC 0.96, and the fifth, the number of reference pools, was a construction "
     "artifact (3 for 22 of 24 attacks, 2 for 137 of 138 market moves).")
h2("3.3 Shortcut reliance causes false alarms when the market changes")
para("Trained on earlier incidents and tested on later ones, the tree ensembles, XGBoost above all, produced "
     "bursts of false alarms (Figure 3). Summed over the seven cutoffs (114 distinct market moves), false alarms were 14 for "
     "the BN, 40 [13, 70] for the random forest, and 99 [61, 141] for XGBoost (means over 30 seeds). Paired against the BN, only XGBoost's excess was "
     "significant (+85 [40, 129]; random forest +26 [\u22128, 63]). Label noise amplifies the bursts: trained on the 9 verified manipulations only, the random "
     "forest raised 0.2 false alarms against the BN's 12, most of its bursts having traced to one bot exploit "
     "measured on the same pool as every market move, while XGBoost still raised 48 (+36 [7, 63] over the BN). Removing the "
     "market-context features cut the tree models' false alarms to 1.2 and 10.5; using only those features produced 102 and 131. The mechanism "
     "is a single drifting feature. Oracle update freshness for market moves rose over time, and it was the only "
     "market-context feature that drifted (all others p \u2265 0.18). The drift coincides with a change in the oracle "
     "itself: it begins when Chainlink launched a new ETH/USD aggregator on 18 September 2024 (median 271 s "
     "before, 576 s after, Mann-Whitney p = 7 \u00d7 10\u207b\u2078), after which updates for moves under "
     "0.5% stopped and heartbeat-only updates rose from 30% to 45%. Removing it alone eliminated all of the "
     "random forest's false alarms and 87% [65, 99] of XGBoost's over 30 seeds. The trees learned cutoffs on this feature, while the "
     "BN's fixed 3,600 s threshold leaves it inert: across attacks, market-context features carried 32% "
     "[23, 44] of the random forest's attribution and 36% [19, 50] of XGBoost's, against 8% [3, 15] for "
     "the BN (Figure 3); permutation importance, computed identically for all three, gave the same order "
     "(0%, 21% and 24% of lost AUC).")
h2("3.4 Verifiable attacks are scarce")
para("Under the verification rule, 12 of the 24 exploits were oracle-price manipulations, and only 9 had "
     "two independent sources; the rest were rounding bugs (3), accounting or access-control flaws (5), "
     "bot or governance exploits (2), or unresolved (2). A second labeler, blind to these labels, applied "
     "the same rule independently: agreement was 15 of 24 (Cohen's κ = 0.46, 95% CI 0.18 to 0.71; 0.50 "
     "for manipulation versus not), and both labelers classified 7 of the 9 verified attacks as "
     "manipulations; restricted to those 7, the market-context shortcut persisted (AUC 0.85 "
     "[0.61, 0.99], permutation p = 0.002). Applying the rule to the 41 Ethereum DeFiHackLabs "
     "incidents a keyword screen flags as price or oracle cases found 20 manipulations, 9 with two sources; "
     "the screen missed 5 of this study's 9, and the union holds 14 verified manipulations. With 9 positives, "
     "the smallest AUC distinguishable from chance (80% power, one-sided \u03b1 = 0.05) is 0.745, against 0.660 with 24. In the pilot, only 3 of the 9 verified attacks were "
     "observable at block resolution; two were restored within the same block, two lacked a reference "
     "price, and two could not yet be read. "
     "Replaying attack blocks transaction by transaction (matching on-chain balances to 10⁻⁵) "
     "recovered both: UwU's source deviated 8.6% mid-block but under 0.1% at every block close, and the "
     "share price Harvest's vault read fell 0.9% within the attack transactions, raising observable "
     "verified attacks to 5 of 9. The "
     "market-context control stayed inconclusive with 5 attacks (one-sided 95% upper bound on its AUC 0.74 "
     "to 0.84); a passing bound below 0.60 would need roughly 40 verified attacks, so no classifier result is claimed.")

# ---------------- 4. Discussion ----------------
h1("4. Discussion")
para("The dominant factor in this problem is how the benchmark is built, not which model is used. When "
     "attacks and market moves are collected and measured differently, context features become proxies "
     "for the label and inflate accuracy, a DeFi instance of shortcut learning [9]. Standard "
     "cross-validation hides this; forward-in-time evaluation exposes it, because shortcut features drift. Model design determines exposure. "
     "The BN stayed stable because its fixed thresholds happen to leave the shortcut features inert. This is reduced exposure, not superior learning: with the shortcut "
     "features removed, both tree models were comparably stable (the random forest also with verified labels), and the BN is not "
     "shortcut-free, since a missing feature leaves it at its prior, so missingness also separates the classes for it (10% of its margin over chance on all 24 attacks, 43% on verified labels).")
para("These results suggest a protocol for valid evaluation: admit both classes through one trigger applied "
     "to the price source; measure every feature on that source, as a price, before the exploit; use an "
     "independent reference price; require that a context-only model score near chance; verify labels with "
     "two sources and a blind second labeler; and freeze the pipeline before testing on new incidents. "
     "Limitations: label agreement is moderate; the main results use the confounded negatives; the pilot "
     "has five attacks; block-level data miss attacks completed within one transaction unless blocks are "
     "replayed; and all data are from Ethereum mainnet. Most DeFiHackLabs incidents occur on other chains, "
     "so multi-chain data and intra-block replay are the most direct paths to an adequately powered test.")

# ---------------- 5. Conclusion ----------------
h1("5. Conclusion")
para("Near-perfect accuracy at telling DeFi oracle attacks from market moves can come from how the data "
     "were collected rather than from the attacks themselves. Here, features unrelated to the attacked "
     "asset predicted the label with an AUC of 0.91, and XGBoost, which leaned on them most, raised bursts of "
     "false alarms when the market changed, while a Bayesian network whose domain-set thresholds ignored them did "
     "not. Only about 14 verifiable manipulations were found on Ethereum mainnet, far fewer than a valid test needs, and some complete "
     "within one block. Progress on this problem depends on benchmarks that source and measure both classes the "
     "same way.")
rich([("Code availability. ", "b"), ("Analysis code and processed features are available from the author on request.", "")])

# ---------------- References ----------------
h1("References")
REFS = [
 "S. Ellis, A. Juels, and S. Nazarov, “ChainLink: A decentralized oracle network,” white paper, 2017.",
 "S. M. Werner, D. Perez, L. Gudgeon, A. Klages-Mundt, D. Harz, and W. J. Knottenbelt, “SoK: Decentralized finance (DeFi),” in Proc. 4th ACM Conf. Advances in Financial Technologies, 2022, pp. 30–46, doi: 10.1145/3558535.3559780.",
 "K. Qin, L. Zhou, B. Livshits, and A. Gervais, “Attacking the DeFi ecosystem with flash loans for fun and profit,” in Financial Cryptography and Data Security, LNCS vol. 12674, 2021, pp. 3–32, doi: 10.1007/978-3-662-64322-8_1.",
 "R. Xi, Z. Wang, and K. Pattabiraman, “POMABuster: Detecting price oracle manipulation attacks in decentralized finance,” in Proc. IEEE Symp. Security and Privacy, 2024, pp. 3923–3942, doi: 10.1109/SP54263.2024.00257.",
 "S. Wu, Z. Yu, D. Wang, Y. Zhou, L. Wu, H. Wang, and X. Yuan, “DeFiRanger: Detecting DeFi price manipulation attacks,” IEEE Trans. Dependable Secure Comput., vol. 21, no. 4, pp. 4147–4161, 2024, doi: 10.1109/TDSC.2023.3346888.",
 "M. Pratiwi and Y.-H. Choi, “DeFiTrace: Event-enriched detection of price oracle manipulation across DeFi transactions,” ACM Trans. Privacy and Security, vol. 29, no. 3, pp. 1–28, 2026, doi: 10.1145/3817054.",
 "D. Venkatraman and M. Kuppusamy, “TOAD-ML: A trust-aware framework for validating oracle price feeds and detecting manipulation in decentralized finance,” Frontiers in Blockchain, vol. 9, Art. no. 1903202, 2026, doi: 10.3389/fbloc.2026.1903202.",
 "S. Chockalingam, W. Pieters, A. Teixeira, and P. van Gelder, “Bayesian network model to distinguish between intentional attacks and accidental technical failures: A case study of floodgates,” Cybersecurity, vol. 4, no. 1, p. 29, 2021, doi: 10.1186/s42400-021-00086-6.",
 "R. Geirhos, J.-H. Jacobsen, C. Michaelis, R. Zemel, W. Brendel, M. Bethge, and F. A. Wichmann, “Shortcut learning in deep neural networks,” Nature Machine Intelligence, vol. 2, no. 11, pp. 665–673, 2020, doi: 10.1038/s42256-020-00257-z.",
 "SunWeb3Sec, “DeFiHackLabs,” GitHub repository. [Online]. Available: https://github.com/SunWeb3Sec/DeFiHackLabs",
 "A. Ankan and A. Panda, “pgmpy: Probabilistic graphical models using Python,” in Proc. 14th Python in Science Conf., 2015, pp. 6–11, doi: 10.25080/Majora-7b98e3ed-001.",
 "T. Chen and C. Guestrin, “XGBoost: A scalable tree boosting system,” in Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining, 2016, pp. 785–794.",
 "S. M. Lundberg and S.-I. Lee, “A unified approach to interpreting model predictions,” in Advances in Neural Information Processing Systems 30, 2017, pp. 4765–4774.",
]
for i, r in enumerate(REFS, 1):
    p = para(f"[{i}] {r}", size=8, align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)

# ---------------- Figure captions ----------------
h1("Figure Captions")
CAPS = [
 ("Figure 1.", "How each class was assembled. Exploits came from public lists and were measured at their attacked venues, mostly after the exploit, with ETH/USD data as a fallback. "
  "Market moves were selected because ETH/USD moved and were measured on one ETH pool."),
 ("Figure 2.", "Features that describe only the Ethereum market (orange) separate the classes as well as "
  "features of the attacked venue (blue), for all 24 labeled attacks and for the 9 verified oracle "
  "manipulations. Bars show random-forest LOOCV AUC with bootstrap 95% CIs; dashed line, chance."),
 ("Figure 3.", "Left: share of each model's attribution carried by market-context features on attack "
  "incidents (bootstrap 95% CIs over 200 refits). Right: forward-in-time false alarms summed "
  "over seven cutoffs (30-seed means; intervals combine seed and incident "
  "resampling), with all, without, and with only market-context features. With market-context features only, the BN flagged nothing. Gray: all features, trained on the 9 verified manipulations only."),
]
for lab, txt in CAPS:
    rich([(lab + " ", "b"), (txt, "")], size=8.5, after=2)

doc.save("/home/claude/cjsj/JillellamudiSanjay_paper.docx")
print("saved")
