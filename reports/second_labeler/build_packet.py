"""Blind second-labeler packet (roadmap A9). Shows ONLY incident name, date, attack tx and
starting sources. Never shows the first labeler's category, S, C, subflags, notes or confidence."""
import csv, random, datetime as dt
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment, PatternFill

SRC = "/mnt/user-data/uploads/oracle-diagnosis/reports/a2_label_sheet.csv"
rows = list(csv.DictReader(open(SRC)))
assert len(rows) == 24
# date from the project's incident list (block -> date is in features.parquet); use parquet
import pandas as pd
feat = pd.read_parquet("/mnt/user-data/uploads/oracle-diagnosis/data/processed/features.parquet")
date_of = dict(zip(feat.incident, feat.date))
random.Random(20260930).shuffle(rows)          # randomized order, fixed seed

items = []
for i, r in enumerate(rows, 1):
    items.append(dict(n=i, incident=r["incident_id"], date=str(date_of.get(r["incident_id"], "")),
                      tx=r["T_tx"], sources=[s.strip() for s in r["sources"].split("|") if s.strip()]))

# ---------------- instructions + rule (docx) ----------------
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
for s in d.sections: s.left_margin = s.right_margin = Inches(0.9)
def P(t="", b=False, size=11, after=4, i=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; r.italic = i; r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after); return p
def B(t):
    p = d.add_paragraph(t, style="List Bullet"); p.paragraph_format.space_after = Pt(2)

P("Second-Labeler Packet: DeFi Exploit Categorization", b=True, size=16, after=6)
P("Thank you for helping. You are an independent second labeler for 24 real DeFi exploits. Your labels "
  "will be compared with another labeler's to measure agreement. Please work blind:", after=4)
B("Do not look at the project repository, the paper, or anyone else's labels.")
B("Do not discuss any incident with Sanjay until you have submitted your answer sheet.")
B("Budget about 5 minutes per incident (about 2 hours total). If you cannot decide, use O.")
P("For each incident, fill one row of the answer sheet (second_labeler_answers.xlsx).", after=8)

P("What to decide", b=True, size=13, after=4)
P("Your job is to decide whether each exploit was an oracle-price manipulation (category M) or something "
  "else. Two terms matter:", after=3)
B("Price source (S): an on-chain value that reports a price or exchange rate, such as a pool's spot price, "
  "a time-weighted average price, a vault's share price, or an exchange rate.")
B("Consumer (C): a different contract that reads S and uses it to value collateral, mint, redeem, swap, "
  "or liquidate.")

P("Category M: oracle-price manipulation. Label M only if ALL four hold:", b=True, after=3)
B("M1. An attacker action (possibly funded by a flash loan) moved the value of a price source S away from "
  "its fair value, at or before the exploit transaction.")
B("M2. A different contract C reads S (directly or through a cached value) and uses it to value "
  "collateral, mint, redeem, swap, or liquidate.")
B("M3. The attacker's profit comes from C relying on the moved value.")
B("M4. S moved because of attacker-controlled flows into S (swaps, deposits, donations, mint or burn, "
  "self-trades), or because the attacker exploited an arithmetic flaw in S's own math while a separate "
  "consumer still read the result. In that second case, also tick the math_enabled column.")

P("If M does not hold, choose the first that fits, in this order:", b=True, after=3)
B("R: a rounding, precision, or invariant-math bug in a pool or curve, with no separate consumer reading "
  "a manipulated price.")
B("A: an accounting, access-control, or logic bug (share accounting, unchecked callbacks, mint bugs) "
  "where the price source and its only consumer are the same contract.")
B("K: an exploit of an MEV bot, a private key, a permission, or governance, unrelated to a consumed price.")
B("O: the evidence is not enough to decide.")
P("Tie-break: if the price source and its only consumer are the same contract and the flaw is in that "
  "contract's own share or rate accounting, label A, even if it looks like manipulation. A donation to "
  "a vault whose rate is read by a different lending contract is M.", i=True, after=8)

P("Evidence rules", b=True, size=13, after=4)
B("Use at least two independent sources per incident when possible: a protocol or auditor post-mortem "
  "plus an independent analysis (for example BlockSec, PeckShield, SlowMist, OpenZeppelin, Halborn, "
  "rekt.news).")
B("The starting links below are suggestions only. You may use any other sources, and you should record "
  "every source you actually relied on.")
B("A proof-of-concept code comment (for example in DeFiHackLabs) alone counts as a single source.")
B("Optional: view the attack transaction on a block explorer (the Etherscan link) to see which contracts "
  "were called.")
P("", after=6)

P("The 24 incidents (randomized order)", b=True, size=13, after=4)
for it in items:
    P(f"{it['n']}. {it['incident']}  ({it['date']})", b=True, after=1)
    P(f"Attack transaction: https://etherscan.io/tx/{it['tx']}", size=9.5, after=1)
    for s in it["sources"]:
        P(f"Starting source: {s}", size=9.5, after=1)
    P("", after=2)
d.save("/home/claude/cjsj/labeler/second_labeler_instructions.docx")

# ---------------- answer sheet (xlsx) ----------------
wb = Workbook(); ws = wb.active; ws.title = "Answers"
hdr = ["#", "Incident", "Date", "Category (M/R/A/K/O)", "math_enabled (Y/N)", "Price source S (what and where)",
       "Consumer C (which contract read S)", "Number of independent sources", "Sources used (URLs)",
       "Confidence (high/medium/low)", "Notes"]
ws.append(hdr)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2A78D6")
    c.alignment = Alignment(wrap_text=True, vertical="top")
for it in items:
    ws.append([it["n"], it["incident"], it["date"], "", "", "", "", "", "", "", ""])
widths = [4, 24, 11, 12, 11, 34, 34, 12, 40, 13, 30]
for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
for dv, col in [(DataValidation(type="list", formula1='"M,R,A,K,O"', allow_blank=True), "D"),
                (DataValidation(type="list", formula1='"Y,N"', allow_blank=True), "E"),
                (DataValidation(type="list", formula1='"high,medium,low"', allow_blank=True), "J")]:
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}25")
dv = DataValidation(type="whole", operator="between", formula1="0", formula2="20", allow_blank=True)
ws.add_data_validation(dv); dv.add("H2:H25")
ws.freeze_panes = "C2"
ws2 = wb.create_sheet("Labeler info")
for r in [["Your name", ""], ["Date completed", ""], ["Approx. hours spent", ""],
          ["Did you see any other labels or the paper before finishing? (Y/N)", ""]]:
    ws2.append(r)
ws2.column_dimensions["A"].width = 62; ws2.column_dimensions["B"].width = 30
wb.save("/home/claude/cjsj/labeler/second_labeler_answers.xlsx")
print("packet built:", len(items), "incidents")
